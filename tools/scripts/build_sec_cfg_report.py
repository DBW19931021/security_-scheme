from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"Z:\code\ngu800\secure\security_-scheme")
OUT = ROOT / "outputs" / "sec-cfg" / "NGU800P_SEC_CFG_寄存器与软件使用专题报告.docx"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F7FB"
GRAY = "5B6573"
RED = "B42318"
AMBER = "9A6700"
GREEN = "16794A"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, code):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    return run


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Title": (25, "202B33", 0, 8),
        "Subtitle": (13, GRAY, 0, 16),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(9)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.05


def configure_section(section):
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("NGU800P D0  |  SEC_CFG专题报告")
    set_run_font(r, size=8.5, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("security_-scheme  ·  内部技术资料  ·  ")
    set_run_font(r, size=8, color=GRAY)
    add_field(fp, "PAGE")


def add_para(doc, text="", bold_prefix=None, color=None, italic=False, align=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, color=color, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, color=color, italic=italic)
    return p


def add_bullets(doc, items, numbered=False):
    for index, item in enumerate(items, start=1):
        if numbered:
            p = doc.add_paragraph()
            prefix = f"{index}. "
        else:
            p = doc.add_paragraph(style="List Bullet")
            prefix = ""
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(prefix + item)
        set_run_font(r)


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    p_pr.append(shd)
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=9)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    keep_row_together(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(value))
        set_run_font(r, size=font_size, bold=True, color=DARK_BLUE)
        if widths:
            cell.width = Inches(widths[idx])
    for row_values in rows:
        row = table.add_row()
        keep_row_together(row)
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            if widths:
                cell.width = Inches(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_status_box(doc, title, text, fill="FFF4E5", color=AMBER):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=160, end=160)
    p = cell.paragraphs[0]
    r = p.add_run(title + "  ")
    set_run_font(r, bold=True, color=color)
    r = p.add_run(text)
    set_run_font(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    set_repeat_keep(p, keep_next=True)
    return p


def build():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)

    # Memo masthead cover.
    add_para(doc, "NGU800P D0 / SOC SECURITY", color=BLUE, after=4)
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("SEC_CFG寄存器与软件使用专题报告")
    set_run_font(r, size=25, bold=True, color="202B33")
    sub = doc.add_paragraph(style="Subtitle")
    r = sub.add_run("架构定位 · 寄存器模型 · 软件接口 · Firewall边界 · DV/Test Case")
    set_run_font(r, size=13, color=GRAY)

    metadata = [
        ("项目", "NGU800P安全方案 / security_-scheme"),
        ("适用", "NGU800P D0"),
        ("版本", "v0.1 · 2026-08-12"),
        ("状态", "REVIEW READY WITH OPEN BINDINGS"),
        ("输入", "SRC-0026（12张内网整理稿截图）+ SRC-0022（RTL同步生成头）"),
        ("关键阻断", "OPEN-CONFLICT-015；关联OPEN-CONFLICT-014"),
    ]
    add_table(doc, ["元数据", "内容"], metadata, widths=[1.25, 5.25], font_size=9.4)
    add_status_box(
        doc,
        "结论先行",
        "截图中的17个offset、status/hardware-error位和拼接顺序可进入参数化设计；截图base/size与当前RTL同步生成头冲突，且baremetal映射仍为无效/地址0，因此不得直接生成产品MMIO常量。",
        fill="FDECEC",
        color=RED,
    )
    add_para(doc, "本报告面向软件安全方案开发、驱动实现评审和DV/Test Case设计。所有测试状态均为NOT_EXECUTED。", italic=True, color=GRAY)
    doc.add_page_break()

    add_heading(doc, "目录与阅读路径", 1)
    toc_p = doc.add_paragraph()
    add_field(toc_p, 'TOC \\o "1-3" \\h \\z \\u')
    add_para(doc, "建议阅读：软件开发关注第4、8、9章；RTL/DV关注第3、5、6、11章；Security Owner关注第7、10、12章。", color=GRAY)

    add_heading(doc, "1. 文档目的、证据状态与结论边界", 1)
    add_para(doc, "本报告将用户提供的12张SEC_CFG内网整理稿截图转录并接入工程工作流，同时与当前RTL同步生成头和baremetal实现状态交叉检查。截图是内部二次整理资料，不等价于原始寄存器Excel、地址表、RTL或目标板Evidence。")
    add_table(doc, ["状态", "本报告中的含义"], [
        ("DOCUMENTED", "截图或当前工程资料明确记录，但尚未证明目标板运行行为"),
        ("CONFIRMED", "由匹配目标build的RTL/EMU/FPGA/硅Evidence验证；本专题当前无此状态"),
        ("PROPOSED", "推荐的软件或验证方法，仍需Owner冻结"),
        ("CONFLICTING", "资料冲突或关键语义缺失，禁止固化为产品事实"),
    ], widths=[1.35, 5.15], font_size=9.2)
    add_bullets(doc, [
        "截图明确记录17个32-bit寄存器；除dbg_en_cfg[1:0]外均为RO。",
        "reset=0只表示寄存器/字段复位值，释放复位后外部镜像可立即更新。",
        "SEC_CFG是状态汇聚与受控观察窗口，不是eHSM完整内部CSR、eFuse编程接口或Firewall配置块。",
        "截图提到的sec_cfg_apb_reg.xlsx、mgmt_addr_map.xlsx、eHSM TRM和芯片安全方案本次未作为原始附件入库。",
    ])

    add_heading(doc, "2. SEC_CFG在系统中的位置", 1)
    add_para(doc, "SEC_CFG位于SoC管理域，汇聚eHSM状态/错误、生命周期、UID和Debug最终输出。Boot MCU、RAS或其他master首先经过Target-side SEC_CFG Firewall，再读取该窗口。")
    add_code(doc, "eHSM/status source ─┐\nLifecycle / UID ─────┼──> SEC_CFG status/observation window ──> Boot SW / GSP / RAS\nDebug policy/auth ───┘                         ▲\nMaster/UserId ──> SEC_CFG Firewall ────────────┘")
    add_table(doc, ["模块", "负责", "不负责"], [
        ("SEC_CFG Firewall", "决定哪个UserId可读/写SEC_CFG；记录拒绝访问", "解释eHSM状态、LCS或Debug输出"),
        ("SEC_CFG", "镜像状态/错误/LCS/UID；配置少量Debug字段", "Firewall policy、eFuse编程、完整eHSM CSR"),
        ("安全软件", "按阶段轮询、保留raw快照、fail-close", "猜测地址、错误bit、LCS或Debug策略"),
        ("RAS/源模块", "错误处置、清除、reset/隔离策略", "通过SEC_CFG只读镜像直接清错"),
    ], widths=[1.25, 2.65, 2.6], font_size=8.7)

    add_heading(doc, "3. 地址、窗口与接口属性", 1)
    add_status_box(doc, "OPEN-CONFLICT-015", "目标D0 SEC_CFG地址、窗口和多项语义未冻结。产品代码只能引用关闭冲突后的权威生成源。", fill="FDECEC", color=RED)
    add_table(doc, ["来源/状态", "Base", "Size", "结论"], [
        ("SRC-0026截图", "0x1010_07B0_0000", "4 KiB", "CONFLICTING；仅转录"),
        ("SRC-0022生成头", "0x1010_0820_0000", "16 KiB", "当前SoC数值权威，但未被status实现绑定"),
        ("baremetal status绑定", "0x0", "—", "MAP_VALID=0；不得实际访问"),
    ], widths=[1.55, 1.75, 1.05, 2.15], font_size=8.8)
    add_bullets(doc, [
        "截图接口：32-bit APB、12-bit local address、pclk、porstn_async。",
        "候选有效寄存器范围为0x000..0x040，共17项。",
        "通用生成项显示set/clear offset 0x800/0x1000，但SEC_CFG字段无set/clear标记，别名地址不可用。",
        "绝对地址、window size和目标instance必须由D0 RTL/CSR生成源及readback Evidence冻结。",
    ])

    add_heading(doc, "4. 完整寄存器表", 1)
    regs = [
        ("000", "hsm_status0", "RO", "0", "o_hsm_status[31:0]"),
        ("004", "hsm_status1", "RO", "0", "o_hsm_status[63:32]"),
        ("008", "hsm_err_hw0", "RO", "0", "o_hsm_err_hw[31:0]"),
        ("00C", "hsm_err_hw1", "RO", "0", "o_hsm_err_hw[63:32]"),
        ("010", "hsm_err_fw0", "RO", "0", "o_hsm_err_fw[31:0]；raw-only"),
        ("014", "hsm_err_fw1", "RO", "0", "o_hsm_err_fw[63:32]；raw-only"),
        ("018", "lcs", "RO", "0", "Life Cycle State；编码待冻结"),
        ("01C", "uid0", "RO", "0", "UID[31:0]"),
        ("020", "uid1", "RO", "0", "UID[63:32]"),
        ("024", "uid2", "RO", "0", "UID[95:64]"),
        ("028", "uid3", "RO", "0", "UID[127:96]"),
        ("02C", "uid4", "RO", "0", "UID[159:128]"),
        ("030", "dbg_en_cfg", "MIXED", "0", "[1:0] RW；[31:2] Reserved RO"),
        ("034", "soc_dbg_en_out0", "RO", "0", "output[31:0]"),
        ("038", "soc_dbg_en_out1", "RO", "0", "output[63:32]"),
        ("03C", "soc_dbg_en_out2", "RO", "0", "output[95:64]"),
        ("040", "soc_dbg_en_out3", "RO", "0", "output[127:96]"),
    ]
    add_table(doc, ["Offset", "Register", "Access", "Reset", "映射/说明"], regs, widths=[0.65, 1.55, 0.7, 0.65, 2.95], font_size=8.2)
    add_para(doc, "注：表内offset是可参数化输入；绝对地址须等待OPEN-CONFLICT-015关闭。", italic=True, color=GRAY)

    add_heading(doc, "5. hsm_status[63:0]与启动判定", 1)
    status_rows = [
        ("63:48", "hsm_fw_sta3", "Firmware状态片段；依匹配FW文档"),
        ("47:32", "hsm_fw_sta2", "Firmware状态片段；依匹配FW文档"),
        ("31:28", "Reserved", "保留"), ("27", "cpu_hart_halted", "eHSM CPU halted"),
        ("26", "Reserved", "保留"), ("25", "cpu_wfi", "eHSM CPU WFI"),
        ("24:21", "hsm_fw_sta1", "Firmware状态片段"),
        ("20/19/18", "soc_cpu_release/reset, soc_reset", "eHSM输出的SoC控制状态"),
        ("17/16", "soc_dbg_en / hsm_dbg_en", "Debug enable状态"),
        ("15", "hsm_fw_sta0", "Firmware状态片段"),
        ("14..8", "undef/destroy/debug/user/manu/dev/test", "生命周期状态bit"),
        ("7/6", "soc_verify_err/done", "SoC镜像验证阶段"),
        ("5/4", "firmware_err/done", "eHSM Firmware阶段"),
        ("3/2", "bootloader_err/done", "eHSM Bootloader阶段"),
        ("1/0", "hw_boot_err/done", "eHSM hardware boot阶段"),
    ]
    add_table(doc, ["Bit", "Field", "含义/边界"], status_rows, widths=[0.85, 2.35, 3.3], font_size=8.3)
    add_heading(doc, "5.1 推荐轮询原则", 2)
    add_bullets(doc, [
        "每次先检查当前阶段error bit，再检查对应done bit。",
        "每个阶段设置有界timeout；error或timeout时采集status/hw error/fw error完整raw值。",
        "hw_boot_done只表示hardware boot完成，不表示Bootloader、Firmware或SoC verify全部ready。",
        "目标启动流程等待到firmware_done还是soc_verify_done由具体stage合同决定。",
        "禁止以‘读到非零’作为ready条件。",
    ], numbered=True)

    add_heading(doc, "6. hsm_err_hw[63:0]硬件错误", 1)
    hw_rows = [
        ("40", "wdt_timeout", "Watchdog timeout"),
        ("35", "otp_key_crc_err", "OTP key CRC error"),
        ("34", "hw_trng_retry_warning", "TRNG健康检测失败并进入retry"),
        ("33", "hw_trng_retry_fail", "TRNG健康检测retry超时"),
        ("32", "hw_trng_ht_fail", "TRNG health test failure"),
        ("29/28/27/26", "soc_err_ahb_cfg/nvm/otp/mem", "AHB response error"),
        ("25/24", "soc_err_axi_dma_rd/wr", "DMA AXI response error"),
        ("19:16", "mem_ecc_mb_pke3..0", "PKE RAM uncorrectable ECC"),
        ("15/14/13/12", "mem_ecc_mb_kmu/dram/iram/irom", "Uncorrectable ECC"),
        ("7:4", "mem_ecc_1b_pke3..0", "PKE RAM correctable ECC"),
        ("3/2/1/0", "mem_ecc_1b_kmu/dram/iram/irom", "Correctable ECC"),
    ]
    add_table(doc, ["Bit", "Field", "含义"], hw_rows, widths=[0.8, 2.75, 3.0], font_size=8.3)
    add_bullets(doc, [
        "Reserved：63:41、39:36、31:30、23:20、11:8。",
        "不存在相应ECC功能时，相关ECC bit按reserved处理。",
        "截图称已定义错误多为level；清除、锁存和重触发行为由源模块/RTL/Firmware合同冻结。",
        "SEC_CFG只有RO镜像，不提供error clear/W1C。",
    ])

    add_heading(doc, "7. Firmware Error、Lifecycle与UID", 1)
    add_heading(doc, "7.1 hsm_err_fw[63:0]", 2)
    add_para(doc, "当前资料没有给出具体bit定义，只说明信号由目标Firmware/Bootloader定义。软件必须记录完整raw 64-bit、镜像build和发生阶段；在匹配版本的BL/FW TRM或error header到位前，不发布逐bit公共枚举、不把某一bit作为最终Expected，也不通过SEC_CFG写操作清错。")
    add_heading(doc, "7.2 lcs", 2)
    add_para(doc, "lcs@0x018是32-bit RO生命周期镜像，但合法编码、one-hot/枚举/eFuse raw属性均未给出。它与hsm_status[14:8]的关系未证明。软件同时读取两者并交叉检查；非法、未定义或不一致时fail-close。")
    add_heading(doc, "7.3 160-bit UID", 2)
    add_code(doc, "uid0 = UID[31:0]\nuid1 = UID[63:32]\nuid2 = UID[95:64]\nuid3 = UID[127:96]\nuid4 = UID[159:128]")
    add_bullets(doc, [
        "接口首选返回uint32_t word[5]，word[0]为低32位。",
        "若输出20-byte或字符串，必须明确word内字节序和显示顺序。",
        "当前没有UID-valid位或稳定时刻；reset后任意时刻读数不能直接作为设备身份PASS。",
    ])

    add_heading(doc, "8. Debug配置与128-bit输出", 1)
    add_table(doc, ["对象", "访问", "当前已知", "未冻结"], [
        ("dbg_en_cfg[1:0]", "RW", "Die1 pre-USER Debug enable候选配置", "bit scope、写Owner、LCS、lock、认证合成"),
        ("dbg_en_cfg[31:2]", "RO", "Reserved", "写响应/副作用"),
        ("soc_dbg_en_out0..3", "RO", "最终128-bit观察输出", "完整bit allocation与consumer gate矩阵"),
    ], widths=[1.55, 0.65, 2.2, 2.1], font_size=8.5)
    add_status_box(doc, "软件门禁", "普通运行软件、Host接口和诊断CLI不得暴露任意dbg_en_cfg写接口。可信启动代码若最终获准写入，只写[1:0]、立即读回，并以最终输出和实际Debug gate Evidence共同验收。", fill="FFF4E5", color=AMBER)
    add_para(doc, "128-bit输出不是2-bit配置的简单复制。没有SoC顶层连接矩阵时，不能根据某个输出bit为1就宣称JTAG、UART或PCIe Debug已经实际开放。")

    add_heading(doc, "9. 软件接口与读取顺序", 1)
    add_heading(doc, "9.1 建议数据结构", 2)
    add_code(doc, "struct sec_cfg_snapshot {\n    uint64_t hsm_status;\n    uint64_t hsm_err_hw;\n    uint64_t hsm_err_fw;\n    uint32_t lcs_raw;\n    uint32_t uid_word[5];\n    uint32_t dbg_cfg_raw;\n    uint32_t dbg_out_word[4];\n};")
    add_heading(doc, "9.2 启动读取顺序", 2)
    add_bullets(doc, [
        "确认eHSM reset已释放。",
        "确认当前master的SEC_CFG Firewall authority已配置。",
        "按目标阶段轮询hsm_status：先error、后done，并设置timeout。",
        "error/timeout时采集完整status、hardware error和firmware error。",
        "读取lcs并与status lifecycle bits交叉检查。",
        "在UID有效阶段读取uid0..uid4并显式序列化。",
        "读取dbg_en_cfg和soc_dbg_en_out0..3；执行最终Debug策略。",
        "不得通过SEC_CFG写操作清除status/error。",
    ], numbered=True)
    add_heading(doc, "9.3 64-bit防撕裂", 2)
    add_code(doc, "do {\n    high_before = mmio_read32(base + high_off);\n    low         = mmio_read32(base + low_off);\n    high_after  = mmio_read32(base + high_off);\n} while (high_before != high_after);\nreturn ((uint64_t)high_after << 32) | low;")
    add_para(doc, "该方法是PROPOSED的best-effort防撕裂，不等于硬件原子快照；跨word同时变化的最终一致性仍需RTL snapshot/latch/CDC合同。")

    add_heading(doc, "10. 与Firewall的关系", 1)
    add_para(doc, "Firewall和SEC_CFG位于不同窗口：Firewall authority决定‘谁能进入SEC_CFG’，SEC_CFG寄存器决定‘进入后能观察或配置什么’。")
    add_table(doc, ["Firewall候选Offset", "Register/信息", "作用"], [
        ("0x004", "F_SECCFG_ENABLE", "check_en[0]、hide_en[1]"),
        ("0x054", "F_SECCFG_AUTHORITY", "每个UserId的R/W authority"),
        ("0x078", "SEC_CFG error clear", "清Firewall错误状态"),
        ("0x07C/0x080", "error address", "被拒绝的48-bit地址"),
        ("0x084", "error info", "write[4]、userid[3:0]"),
        ("0x088", "error status", "错误状态"),
    ], widths=[1.25, 2.2, 3.05], font_size=8.5)
    add_bullets(doc, [
        "这些CSR属于Firewall，不属于SEC_CFG自身。",
        "hide模式返回OK不代表安全成功，必须证明Slave无访问/副作用且错误快照正确。",
        "Firewall候选base和完整CSR绑定继续由OPEN-CONFLICT-014管理。",
    ])

    add_heading(doc, "11. DV与Test Case建议", 1)
    cases = [
        ("SECCFG-REG-001", "地址与基本读", "BLOCKED", "OPEN-CONFLICT-015"),
        ("SECCFG-STATUS-001", "启动阶段/error-done-timeout", "READY_FOR_MODEL", "—"),
        ("SECCFG-HWERR-001", "ECC/TRNG/bus/watchdog注入", "READY_FOR_DV", "—"),
        ("SECCFG-FWERR-001", "BL/FW故障与raw值", "BLOCKED", "OPEN-CONFLICT-015"),
        ("SECCFG-LCS-001", "生命周期与非法值", "BLOCKED", "OPEN-CONFLICT-015"),
        ("SECCFG-UID-001", "160-bit稳定性/端序", "BLOCKED", "OPEN-CONFLICT-015"),
        ("SECCFG-DBG-001", "cfg/output/实际gate", "BLOCKED", "OPEN-CONFLICT-015"),
        ("SECCFG-FW-001", "未授权UserId访问", "BLOCKED", "OPEN-CONFLICT-014"),
        ("SECCFG-RO-001", "RO offset写行为", "READY_FOR_DV", "—"),
        ("SECCFG-RESET-001", "reset与动态镜像", "READY_FOR_DV", "—"),
        ("SECCFG-ATOMIC-001", "跨word一致性", "BLOCKED", "OPEN-CONFLICT-015"),
        ("SECCFG-ORDER-001", "阶段顺序防误判", "READY_FOR_MODEL", "—"),
    ]
    add_table(doc, ["Test ID", "场景", "设计就绪", "阻断"], cases, widths=[1.55, 2.65, 1.35, 1.55], font_size=8.0)
    add_para(doc, "全部12条用例均为NOT_EXECUTED。READY_FOR_MODEL/READY_FOR_DV只表示具备开始实现测试的输入，不表示目标硬件已通过。", bold_prefix="全部12条用例均为NOT_EXECUTED。")
    add_heading(doc, "11.1 Evidence最小集合", 2)
    add_bullets(doc, [
        "DUT、RTL、eHSM BL/FW、BootROM/GSP build和执行平台。",
        "当前UserId、LCS、reset类型、Firewall配置和注入方法。",
        "原始SEC_CFG/Firewall寄存器dump、日志和时间点。",
        "Expected的Source/版本、PASS/FAIL/INCONCLUSIVE结论和Owner。",
        "环境不具备注入或观测能力时标记BLOCKED，不得标记PASS。",
    ])

    add_heading(doc, "12. 编码门禁与未关闭项", 1)
    gates = [
        ("Base/size/instance", "RTL/Integration", "SRC-0022或后继生成源 + readback", "阻断产品MMIO"),
        ("FW error bit", "eHSM BL/FW", "匹配image build的TRM/error header", "仅raw记录"),
        ("LCS编码", "Security/eFuse", "唯一编码表及非法值策略", "阻断LCS API"),
        ("Debug规则", "Security/RTL", "scope、Owner、LCS、lock、认证与128-bit矩阵", "阻断普通写API"),
        ("UID-valid", "RTL/Platform", "复位/启动稳定时刻Evidence", "阻断最终身份PASS"),
        ("64-bit一致性", "RTL/DV", "snapshot/latch/CDC合同", "best-effort限定"),
    ]
    add_table(doc, ["待冻结项", "Owner", "关闭证据", "当前影响"], gates, widths=[1.35, 1.25, 2.8, 1.6], font_size=8.2)
    add_status_box(doc, "禁止事项", "不得把截图base、地址0占位、未定义FW error/LCS/Debug语义或set/clear别名写入产品代码；不得把未执行测试表述为PASS。", fill="FDECEC", color=RED)

    add_heading(doc, "13. 工程落盘与追踪", 1)
    artifacts = [
        ("输入归档", "source-vault/internal-specs/sec-cfg/SRC-0026/"),
        ("来源卡/摄取", "sources/source-cards/SRC-0026.md；sources/intake-reports/SRC-0026-sec-cfg-intake.md"),
        ("冲突", "sources/conflict-reports/CONFLICT-SRC-0026-SEC-CFG-ADDRESS-AND-SEMANTICS.md"),
        ("机读模型", "requirements/sec-cfg-register-map.yaml；requirements/sec-cfg-requirements.yaml"),
        ("架构/接口", "docs/03-architecture/sec-cfg-status-observation.md；docs/04-interfaces/sec-cfg-register-interface.md"),
        ("主详设", "docs/05-software-design/NGU800P安全软件详细设计.md §5.9.2"),
        ("测试", "docs/06-verification/sec-cfg-test-design.md；tests/cases/sec-cfg/；tests/matrices/sec-cfg-traceability.yaml"),
    ]
    add_table(doc, ["层级", "工程路径"], artifacts, widths=[1.25, 5.25], font_size=8.5)

    add_heading(doc, "附录A：来源与截图清单", 1)
    add_para(doc, "SRC-0026共归档12张PNG；第9、10张SHA-256一致，作为原始输入重复件保留。完整哈希见转录文件和source-index.yaml。")
    images = [
        ("01", "purpose-and-sources", "目的、边界、资料基线"),
        ("02", "address-and-boundary", "地址、接口属性、功能边界"),
        ("03", "register-map-part1", "完整寄存器表前半"),
        ("04", "register-map-and-status-part1", "寄存器表后半、status前半"),
        ("05", "status-part2", "status后半、轮询原则"),
        ("06", "hardware-error-part1", "hardware error前半"),
        ("07", "hardware-error-and-fw-error", "hardware error后半、FW error"),
        ("08", "lcs-uid-debug-part1", "LCS、UID、Debug前半"),
        ("09/10", "debug-firewall", "Debug后半、Firewall关系；重复件"),
        ("11", "software-read-guidance", "软件读取宏与64-bit方法"),
        ("12", "dv-tests-and-gaps", "DV建议与未关闭项"),
    ]
    add_table(doc, ["截图", "归档名", "主题"], images, widths=[0.75, 2.35, 3.4], font_size=8.5)

    add_heading(doc, "附录B：审阅结论", 1)
    add_bullets(doc, [
        "可直接进入编码设计：寄存器offset、raw snapshot结构、已列status/hardware-error bit、UID word拼接、error-first轮询、参数化fake-MMIO测试。",
        "需在产品编码前关闭：base/size、FW error位、LCS编码、Debug写策略、UID-valid、64-bit一致性。",
        "需与Firewall专题联合关闭：访问authority、hide响应安全判定及Firewall错误诊断CSR。",
        "当前工程资料可指导test case设计，但所有条目仍为NOT_EXECUTED。",
    ])

    # Core properties and save.
    doc.core_properties.title = "NGU800P D0 SEC_CFG寄存器与软件使用专题报告"
    doc.core_properties.subject = "SEC_CFG architecture, register model, software interface and verification"
    doc.core_properties.author = "NGU800P Security Scheme Team"
    doc.core_properties.keywords = "NGU800P, SEC_CFG, eHSM, Firewall, lifecycle, UID, debug, DV"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
